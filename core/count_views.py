import csv
import uuid
from io import BytesIO, StringIO

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from . import counts
from . import services as s
from .models import Product, StockCount
from .views import protected, problem


def _count_analysis(count):
    lines = list(count.lines.select_related("product"))
    variance_units = sum(abs(line.variance or 0) for line in lines if line.counted is not None)
    variance_value = sum(
        (abs(line.variance or 0) * line.product.cost for line in lines if line.counted is not None),
        0,
    )
    matched = 0
    exceptions = 0
    shortages = 0
    excesses = 0
    for line in lines:
        line.variance_value = abs(line.variance or 0) * line.product.cost
        line.risk = "high" if line.variance_value >= 1000 else ("medium" if line.variance_value >= 250 else "low")
        if line.counted is None:
            line.result_state = "not_counted"
            line.result_label = "Not counted"
        elif line.variance == 0:
            line.result_state = "matched"
            line.result_label = "Matches system"
            matched += 1
        elif line.variance < 0:
            line.result_state = "short"
            line.result_label = "Short"
            shortages += 1
            exceptions += 1
        else:
            line.result_state = "over"
            line.result_label = "Over"
            excesses += 1
            exceptions += 1
    return {
        "lines": lines,
        "variance_units": variance_units,
        "variance_value": variance_value,
        "exception_count": exceptions,
        "matched_count": matched,
        "shortage_count": shortages,
        "excess_count": excesses,
    }


@protected("operate_inventory|approve_operations|manage_company")
def index(request, branch):
    if request.method == "POST":
        count = counts.start_count(request.user, branch, request.POST.get("key"), request.POST.get("category", ""))
        return redirect("stock_count", pk=count.pk)
    return render(request, "counts.html", {
        "title": "Physical stock counts", "key": uuid.uuid4(),
        "categories": Product.objects.filter(active=True).exclude(category="").order_by("category").values_list("category", flat=True).distinct(),
        "rows": StockCount.objects.filter(branch=branch).select_related("created_by", "reviewed_by")[:100],
    })


@protected("operate_inventory|approve_operations|manage_company")
def detail(request, branch, pk):
    count = get_object_or_404(StockCount, pk=pk, branch=branch)
    lines = list(count.lines.select_related("product"))
    status = 200
    if request.method == "POST":
        action = request.POST.get("action")
        try:
            if action in ("save", "submit"):
                values = {
                    str(line.pk): (request.POST.get("quantity_" + str(line.pk), ""), "")
                    for line in lines
                }
                counts.save_count(request.user, branch, pk, values, request.POST.get("note", ""), action == "submit")
                owner_direct = request.POST.get("owner_direct") == "1" and (
                    request.user.is_superuser or request.user.has_perm("core.manage_company")
                )
                if action == "submit" and owner_direct:
                    counts.review_count(
                        request.user, branch, pk, "approve",
                        "Owner direct authority — inventory verification posted after complete blind count.",
                        owner_direct=True,
                    )
            else:
                owner_direct = request.POST.get("owner_direct") == "1" and (
                    request.user.is_superuser or request.user.has_perm("core.manage_company")
                )
                counts.review_count(
                    request.user, branch, pk, action, request.POST.get("review_note", ""),
                    owner_direct=owner_direct,
                )
            messages.success(request, "Stock count recorded.")
            return redirect("stock_count", pk=pk)
        except ValidationError as exc:
            status = 400
            messages.error(request, problem(exc))
            count.refresh_from_db()
            if count.status == "draft" and count.created_by_id == request.user.pk and action in ("save", "submit"):
                for line in lines:
                    line.counted = request.POST.get("quantity_" + str(line.pk), "")
                count.note = request.POST.get("note", "")
    count.refresh_from_db()
    analysis = _count_analysis(count)
    owner = request.user.is_superuser or request.user.has_perm("core.manage_company")
    return render(request, "count.html", {
        "title": "Inventory verification",
        "count": count,
        "editable": count.status == "draft" and count.created_by_id == request.user.pk,
        "can_review": count.status == "submitted"
            and (request.user.has_perm("core.approve_operations") or owner)
            and (count.created_by_id != request.user.pk or owner),
        "owner_direct": owner,
        **analysis,
    }, status=status)


def _export_rows(count):
    analysis = _count_analysis(count)
    rows = []
    for line in analysis["lines"]:
        rows.append([
            line.product.sku,
            line.product.name,
            line.product.category or "",
            line.product.base_unit,
            line.counted if line.counted is not None else "",
            line.expected,
            line.variance if line.variance is not None else "",
            line.result_label,
            line.risk.title(),
            line.variance_value,
        ])
    headers = ["SKU", "Product", "Category", "Unit", "Physical count", "System count", "Variance", "Result", "Risk", "Value exposure"]
    return analysis, headers, rows


@protected("operate_inventory|approve_operations|manage_company")
def export_count(request, branch, pk, format):
    count = get_object_or_404(
        StockCount.objects.select_related("created_by", "reviewed_by"),
        pk=pk, branch=branch,
    )
    if count.status == "draft":
        raise ValidationError("Submit the blind count before exporting its comparison results.")
    analysis, headers, rows = _export_rows(count)
    currency = s.company_policy().currency
    filename = f"stock-count-{str(count.pk)[:8]}"

    if format == "csv":
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["KOFAD Inventory Verification"])
        writer.writerow(["Scope", count.scope or "All active products"])
        writer.writerow(["Status", count.status.title()])
        writer.writerow(["Counter", count.created_by.get_username()])
        writer.writerow(["Matched", analysis["matched_count"]])
        writer.writerow(["Exceptions", analysis["exception_count"]])
        writer.writerow(["Absolute variance units", analysis["variance_units"]])
        writer.writerow(["Value exposure", f"{currency} {analysis['variance_value']:.2f}"])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(buffer.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        return response

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Verification Results"
        sheet.merge_cells("A1:J1")
        sheet["A1"] = "KOFAD Inventory Verification"
        sheet["A1"].font = Font(size=18, bold=True)
        sheet["A2"] = "Scope"
        sheet["B2"] = count.scope or "All active products"
        sheet["A3"] = "Status"
        sheet["B3"] = count.status.title()
        sheet["D2"] = "Matched"
        sheet["E2"] = analysis["matched_count"]
        sheet["D3"] = "Exceptions"
        sheet["E3"] = analysis["exception_count"]
        sheet["G2"] = "Variance units"
        sheet["H2"] = analysis["variance_units"]
        sheet["G3"] = "Value exposure"
        sheet["H3"] = float(analysis["variance_value"])
        for cell in sheet[5]:
            pass
        header_row = 5
        for idx, value in enumerate(headers, 1):
            cell = sheet.cell(header_row, idx, value)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCEEF0")
            cell.alignment = Alignment(horizontal="center")
        for r_idx, row in enumerate(rows, header_row + 1):
            for c_idx, value in enumerate(row, 1):
                sheet.cell(r_idx, c_idx, value)
        widths = [16, 30, 18, 12, 16, 16, 12, 18, 12, 18]
        for idx, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:J{sheet.max_row}"
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response

    if format == "docx":
        from docx import Document as WordDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = WordDocument()
        title = doc.add_heading("KOFAD Inventory Verification", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f"Scope: {count.scope or 'All active products'}  |  Status: {count.status.title()}  |  "
            f"Counter: {count.created_by.get_username()}"
        )
        doc.add_paragraph(
            f"Matched: {analysis['matched_count']}  |  Exceptions: {analysis['exception_count']}  |  "
            f"Absolute variance: {analysis['variance_units']} units  |  "
            f"Value exposure: {currency} {analysis['variance_value']:.2f}"
        )
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for idx, value in enumerate(headers):
            table.rows[0].cells[idx].text = value
        for row in rows:
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                cells[idx].text = str(value)
        if count.note:
            doc.add_heading("Verification note", level=2)
            doc.add_paragraph(count.note)
        output = BytesIO()
        doc.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.docx"'
        return response

    if format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = BytesIO()
        pdf = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        )
        styles = getSampleStyleSheet()
        story = [
            Paragraph("KOFAD Inventory Verification", styles["Title"]),
            Paragraph(
                f"Scope: {count.scope or 'All active products'} &nbsp;&nbsp; Status: {count.status.title()} "
                f"&nbsp;&nbsp; Counter: {count.created_by.get_username()}",
                styles["BodyText"],
            ),
            Paragraph(
                f"Matched: {analysis['matched_count']} &nbsp;&nbsp; Exceptions: {analysis['exception_count']} "
                f"&nbsp;&nbsp; Absolute variance: {analysis['variance_units']} units "
                f"&nbsp;&nbsp; Value exposure: {currency} {analysis['variance_value']:.2f}",
                styles["BodyText"],
            ),
            Spacer(1, 6 * mm),
        ]
        data = [headers] + [[str(value) for value in row] for row in rows]
        table = Table(data, repeatRows=1, colWidths=[23*mm, 43*mm, 28*mm, 17*mm, 24*mm, 24*mm, 18*mm, 24*mm, 18*mm, 25*mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEEF0")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0A2B43")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#C9D6DC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFB")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        if count.note:
            story += [Spacer(1, 5 * mm), Paragraph(f"<b>Verification note:</b> {count.note}", styles["BodyText"])]
        pdf.build(story)
        response = HttpResponse(output.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response

    raise ValidationError("Choose PDF, Excel, Word or CSV.")
