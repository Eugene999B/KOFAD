import os
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import Group, Permission, User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.db import close_old_connections, connection, connections, transaction, DatabaseError
from django.test import TestCase, TransactionTestCase
from django.utils import timezone

from . import services as s
from .models import Access, Audit, Branch, Company, Document, Line, Movement, Party, Payment, Product, Stock


class Fixtures:
    def setup_data(self):
        self.user = User.objects.create_superuser("owner", "owner@example.test", "test-password-long-enough")
        self.reviewer = User.objects.create_superuser("reviewer", "reviewer@example.test", "test-password-long-enough")
        self.branch = Branch.objects.create(name="Main", code="main")
        self.other = Branch.objects.create(name="Warehouse", code="warehouse")
        Company.objects.create()
        self.product = Product.objects.create(name="Test carton", sku="TEST", pack_size=12, retail_unit="50",
            retail_pack="550", wholesale_pack="480", cost="20")
        Stock.objects.create(branch=self.branch, product=self.product, quantity=240)
        self.customer = Party.objects.create(branch=self.branch, kind="customer", name="Customer", phone="test", credit_limit=5000)
        self.supplier = Party.objects.create(branch=self.branch, kind="supplier", name="Supplier", phone="test")

    def sale(self, quantity=1, mode="retail_unit", payments=None, **kwargs):
        payload = {"items":[{"product":self.product.pk, "mode":mode, "quantity":quantity}],
                   "payments":payments if payments is not None else [{"method":"cash", "amount":str(Decimal(50) * Decimal(str(quantity))) }],
                   "party": self.customer.pk, **kwargs}
        return s.post_trade(self.user, self.branch, payload, uuid.uuid4())

    def authenticate_client(self, user=None):
        user = user or self.user
        self.client.force_login(user)
        user.access.refresh_from_db()
        session = self.client.session
        session["access_version"] = user.access.session_version
        session["mfa_ok"] = True
        session["branch"] = self.branch.pk
        session.save()


class BusinessTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def test_pack_math(self):
        self.sale(5)
        self.sale(3, "wholesale_pack", [{"method":"cash","amount":"1440"}])
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, 199)

    def test_disabled_mode_and_negative_quantity(self):
        for quantity, mode in [(1,"wholesale_unit"), (-1,"retail_unit"), (1.5,"retail_unit")]:
            with self.assertRaises(ValidationError):
                self.sale(quantity, mode)
        self.assertEqual(Document.objects.count(), 0)
        self.assertEqual(Movement.objects.count(), 0)

    def test_oversell_rolls_back_everything(self):
        with self.assertRaises(ValidationError):
            self.sale(241)
        self.assertEqual(Document.objects.count(), 0)
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, 240)

    def test_idempotency_and_mismatched_replay(self):
        key = uuid.uuid4()
        payload = {"items":[{"product":self.product.pk,"mode":"retail_unit","quantity":1}], "payments":[{"method":"cash","amount":"50"}], "party": self.customer.pk}
        one = s.post_trade(self.user,self.branch,payload,key)
        two = s.post_trade(self.user,self.branch,payload,key)
        self.assertEqual(one.pk,two.pk)
        self.assertEqual(Movement.objects.count(),1)
        payload["items"][0]["quantity"] = 2
        with self.assertRaises(ValidationError):
            s.post_trade(self.user,self.branch,payload,key)

    def test_credit_requires_customer_due_date_and_limit(self):
        with self.assertRaises(ValidationError):
            self.sale(payments=[], party=None)
        with self.assertRaises(ValidationError):
            self.sale(payments=[],party=self.customer.pk)
        self.customer.credit_limit = 20
        self.customer.save()
        with self.assertRaises(ValidationError):
            self.sale(payments=[],party=self.customer.pk,due_date=timezone.localdate().isoformat())

    def test_credit_collection_and_overallocation(self):
        invoice = self.sale(2,payments=[{"method":"cash","amount":"20"}],party=self.customer.pk,due_date=timezone.localdate().isoformat())
        payment = s.post_payment(self.user,self.branch,{"invoice":str(invoice.pk),"amount":"30","method":"momo"},uuid.uuid4())
        self.assertEqual(s.balance(invoice),Decimal("50"))
        with self.assertRaises(ValidationError):
            s.post_payment(self.user,self.branch,{"invoice":str(invoice.pk),"amount":"51","method":"cash"},uuid.uuid4())
        self.assertEqual(payment.allocations.count(),1)

    def test_return_credit_then_refund_and_quantity_limit(self):
        invoice = self.sale(2,payments=[{"method":"cash","amount":"20"}],party=self.customer.pk,due_date=timezone.localdate().isoformat())
        returned = s.post_return(self.user,self.branch,{"line":invoice.lines.get().pk,"quantity":2,"reason":"Customer returned goods","method":"cash"},uuid.uuid4())
        self.assertEqual(returned.paid,Decimal("20"))
        self.assertEqual(s.balance(invoice),Decimal("0"))
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,240)
        with self.assertRaises(ValidationError):
            s.post_return(self.user,self.branch,{"line":invoice.lines.get().pk,"quantity":1,"reason":"Duplicate return","method":"cash"},uuid.uuid4())

    def test_mixed_payments_and_overpayment(self):
        sale = self.sale(2,payments=[{"method":"cash","amount":"25"},{"method":"momo","amount":"75"}])
        self.assertEqual(sale.payments.count(),2)
        with self.assertRaises(ValidationError):
            self.sale(payments=[{"method":"cash","amount":"51"}])

    def test_purchase_pack_receipt(self):
        doc = s.post_trade(self.user,self.branch,{"party":self.supplier.pk,"items":[{"product":self.product.pk,
            "mode":"retail_pack","quantity":100,"price":"240"}],"payments":[{"method":"bank","amount":"24000"}]},uuid.uuid4(),"purchase")
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,1440)
        self.assertEqual(doc.payments.get().direction,-1)

    def test_transfer_approval_dispatch_receive(self):
        op = s.request_operation(self.user,self.branch,{"kind":"transfer","product":self.product.pk,
            "quantity":12,"destination":self.other.pk,"reason":"Replenish warehouse"})
        with self.assertRaises(ValidationError):
            s.advance_operation(self.user,op.pk,"approve")
        s.advance_operation(self.reviewer,op.pk,"approve")
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,240)
        s.advance_operation(self.user,op.pk,"dispatch")
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,228)
        self.assertFalse(Stock.objects.filter(branch=self.other,product=self.product).exists())
        s.advance_operation(self.user,op.pk,"receive")
        self.assertEqual(Stock.objects.get(branch=self.other,product=self.product).quantity,12)
        s.advance_operation(self.user,op.pk,"receive")
        self.assertEqual(Stock.objects.get(branch=self.other,product=self.product).quantity,12)

    def test_closing_and_independent_verification(self):
        self.sale()
        closing = s.submit_closing(self.user,self.branch,timezone.localdate(),{"cash":"50"},"")
        with self.assertRaises(ValidationError):
            self.sale()
        with self.assertRaises(ValidationError):
            s.verify_closing(self.user,closing)
        s.verify_closing(self.reviewer,closing)
        closing.refresh_from_db()
        self.assertEqual(closing.verified_by,self.reviewer)

    def test_negative_channel_net_can_be_reconciled(self):
        s.post_expense(self.user,self.branch,{"amount":"20","method":"cash","note":"Transport expense"},uuid.uuid4())
        closing = s.submit_closing(self.user,self.branch,timezone.localdate(),{"cash":"-20"},"")
        self.assertEqual(closing.expected["cash"],"-20.00")

    def test_branch_scope_and_permissions(self):
        cashier = User.objects.create_user("cashier",password="cashier-long-password")
        cashier.user_permissions.add(Permission.objects.get(codename="operate_sales"))
        cashier.access.branches.add(self.other)
        with self.assertRaises(PermissionDenied):
            s.permit(cashier,self.branch,"operate_sales")
        with self.assertRaises(PermissionDenied):
            s.permit(cashier,self.other,"operate_finance")

    def test_money_rejects_nan_infinity_and_fractional_cents(self):
        for value in ("NaN","Infinity","-1","1.001","999999999999999999999"):
            with self.assertRaises(ValidationError):
                s.money(value)

    def test_session_revoked_on_security_change(self):
        self.authenticate_client()
        self.assertEqual(self.client.get("/").status_code,200)
        self.user.is_active = False
        self.user.save()
        self.assertEqual(self.client.get("/").status_code,302)

    def test_all_pages_render(self):
        self.authenticate_client()
        for path in ["/","/inventory/","/sales/new/","/purchasing/","/documents/","/parties/",
                     "/finance/","/creditors/","/accounting/","/payroll/","/payroll/rules/","/workers/","/returns/","/operations/","/closings/","/reports/","/audit/","/settings/","/settings/company/",
                     "/administration/","/administration/users/","/administration/roles/","/exports/","/communications/",
                     "/products/new/","/parties/new/"]:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code,200)
        doc = self.sale()
        self.assertEqual(self.client.get(f"/documents/{doc.pk}/").status_code,200)
        self.assertEqual(self.client.get(f"/parties/{self.customer.pk}/statement/").status_code,200)

    def test_exports_are_valid_files(self):
        self.sale()
        self.authenticate_client()
        from openpyxl import load_workbook
        from docx import Document as WordDocument
        for format in ("pdf","xlsx","docx","csv"):
            response = self.client.get(f"/reports/export/{format}/")
            self.assertEqual(response.status_code,200)
            if format == "pdf":
                self.assertTrue(response.content.startswith(b"%PDF"))
            if format == "xlsx":
                sheet = load_workbook(BytesIO(response.content)).active
                self.assertEqual(sheet["A1"].value, Company.objects.first().name)
                self.assertIn("Reference", [cell.value for cell in sheet["A"]])
                self.assertIsNotNone(sheet.freeze_panes)
            if format == "docx":
                document = WordDocument(BytesIO(response.content))
                self.assertGreaterEqual(len(document.tables), 1)
                self.assertIn(Company.objects.first().name, "\n".join(p.text for p in document.paragraphs))

    def test_csrf_is_enforced(self):
        from django.test import Client
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.post("/login/",{"username":"owner","password":"test-password-long-enough"}).status_code,403)

    def test_login_rate_limit(self):
        for _ in range(5):
            self.client.post("/login/",{"username":"owner","password":"wrong"})
        response = self.client.post("/login/",{"username":"owner","password":"test-password-long-enough"})
        self.assertContains(response,"Too many attempts")

    def test_login_session_survives_last_login_update(self):
        response = self.client.post("/login/",{"username":"owner","password":"test-password-long-enough"})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.client.get("/mfa/").url,"/")
        self.assertEqual(self.client.get("/").status_code,200)

    def test_login_accepts_ghana_phone_and_forces_temporary_password_change(self):
        self.user.access.recovery_phone = "+233241234567"
        self.user.access.force_password_change = True
        self.user.access.save(update_fields=["recovery_phone", "force_password_change"])
        response = self.client.post("/login/", {
            "username": "0241234567",
            "password": "test-password-long-enough",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/account/password/")
        self.assertEqual(self.client.get("/").url, "/account/password/")

    def test_owner_provision_command_creates_equal_full_admin_without_storing_password_in_code(self):
        with patch.dict(os.environ, {
            "KOFAD_OWNER_ADMIN_NAME": "Test Owner",
            "KOFAD_OWNER_ADMIN_PHONE": "0249998877",
            "KOFAD_OWNER_ADMIN_INITIAL_PASSWORD": "0249998877",
        }, clear=False):
            call_command("provision_owner_admin", confirm_owner_admin=True)
        owner = User.objects.get(username="0249998877")
        self.assertEqual(owner.get_full_name(), "Test Owner")
        self.assertTrue(owner.is_superuser)
        self.assertTrue(owner.is_staff)
        self.assertTrue(owner.check_password("0249998877"))
        self.assertEqual(owner.access.recovery_phone, "+233249998877")
        self.assertTrue(owner.access.force_password_change)
        self.assertEqual(
            set(owner.access.branches.values_list("pk", flat=True)),
            set(Branch.objects.filter(active=True).values_list("pk", flat=True)),
        )

    def test_malformed_report_range_returns_400(self):
        self.authenticate_client()
        self.assertEqual(self.client.get("/reports/?start=not-a-date").status_code,400)


class ConcurrencyTests(Fixtures, TransactionTestCase):
    def setUp(self):
        self.setup_data()
        Stock.objects.filter(branch=self.branch,product=self.product).update(quantity=1)

    def test_two_buyers_cannot_oversell_last_unit(self):
        def attempt(_):
            close_old_connections()
            try:
                user = User.objects.get(pk=self.user.pk)
                branch = Branch.objects.get(pk=self.branch.pk)
                s.post_trade(user,branch,{"items":[{"product":self.product.pk,"mode":"retail_unit","quantity":1}],
                    "payments":[{"method":"cash","amount":"50"}], "party": self.customer.pk},uuid.uuid4())
                return "posted"
            except ValidationError:
                return "blocked"
            finally:
                connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(attempt,range(2)))
        self.assertCountEqual(results,["posted","blocked"])
        self.assertEqual(Document.objects.count(),1)
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,0)


class LedgerIntegrityTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def test_posted_document_cannot_be_rewritten(self):
        doc = self.sale()
        with self.assertRaises(DatabaseError), transaction.atomic():
            Document.objects.filter(pk=doc.pk).update(total=0)

    def test_ledger_entries_cannot_be_deleted(self):
        self.sale()
        for model in (Line, Payment, Movement, Audit):
            with self.subTest(model=model.__name__):
                with self.assertRaises(DatabaseError), transaction.atomic():
                    model.objects.all().delete()

    def test_signed_closing_cannot_be_recounted(self):
        self.sale()
        closing = s.submit_closing(self.user,self.branch,timezone.localdate(),{"cash":"50"},"")
        with self.assertRaises(DatabaseError), transaction.atomic():
            type(closing).objects.filter(pk=closing.pk).update(counted={"cash":"0"})


class WorkforcePayrollAccountingTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def worker(self):
        from .models import Worker
        return Worker.objects.create(
            branch=self.branch, employee_code="KFD-0001", first_name="Ama", last_name="Mensah",
            phone="0240000000", department="Operations", job_title="Store Officer",
            employment_type="permanent", status="active", hire_date="2026-01-01",
            base_salary=Decimal("1000.00"), recurring_allowance=Decimal("0.00"),
            ghana_card_number="GHA-TEST", ssnit_number="SSNIT-TEST",
            bank_name="Test Bank", bank_account_number="123456",
            created_by=self.user,
        )

    def test_current_ghana_payroll_math_and_control_workflow(self):
        from . import payroll_engine
        from .models import PayrollEntry, PayrollRule
        worker = self.worker()
        rule = PayrollRule.objects.filter(effective_from__lte="2026-10-31").order_by("-effective_from").first()
        self.assertIsNotNone(rule)
        self.assertEqual(rule.employee_ssnit_rate, Decimal("5.5000"))
        period = payroll_engine.create_period(self.user, self.branch, 2026, 10)
        entry = PayrollEntry.objects.get(period=period, worker=worker)
        self.assertEqual(entry.ssnit_employee, Decimal("55.00"))
        self.assertEqual(entry.paye_tax, Decimal("44.98"))
        self.assertEqual(entry.net_pay, Decimal("900.02"))
        period, issues = payroll_engine.prepare_period(self.user, period)
        self.assertFalse(any("negative" in issue.lower() for issue in issues))
        with self.assertRaises(ValidationError):
            payroll_engine.approve_period(self.user, period)
        period = payroll_engine.approve_period(self.reviewer, period)
        period = payroll_engine.lock_period(self.reviewer, period)
        payment = payroll_engine.record_payment(
            self.user, entry, entry.net_pay, "bank", "TEST-PAY-001", "October salary"
        )
        self.assertEqual(payment.amount, Decimal("900.02"))
        period, outstanding = payroll_engine.reconcile_period(period)
        self.assertFalse(outstanding)
        self.assertEqual(period.status, "reconciled")

    def test_worker_private_document_id_card_and_filtered_exports(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import WorkerDocument
        worker = self.worker()
        self.authenticate_client()
        response = self.client.post(f"/workers/{worker.pk}/documents/", {
            "category": "contract",
            "title": "Employment contract",
            "document_number": "CON-001",
            "file": SimpleUploadedFile("contract.pdf", b"%PDF-1.4 KOFAD TEST", content_type="application/pdf"),
        })
        self.assertEqual(response.status_code, 302)
        document = WorkerDocument.objects.get(worker=worker)
        self.assertEqual(len(document.checksum_sha256), 64)
        downloaded = self.client.get(f"/workers/{worker.pk}/documents/{document.pk}/")
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, b"%PDF-1.4 KOFAD TEST")
        card = self.client.get(f"/workers/{worker.pk}/id-card.pdf")
        self.assertEqual(card.status_code, 200)
        self.assertTrue(card.content.startswith(b"%PDF"))
        workforce = self.client.get("/workers/export/xlsx/?department=Operations&joined_from=2026-01-01")
        self.assertEqual(workforce.status_code, 200)

    def test_accounting_reports_and_payroll_pages_share_source_records(self):
        from . import payroll_engine
        worker = self.worker()
        self.sale(2)
        expense = s.post_expense(
            self.user, self.branch,
            {"amount": "20", "method": "cash", "note": "Fuel for local delivery", "category": "fuel"},
            uuid.uuid4(),
        )
        self.assertEqual(expense.expense_category, "fuel")
        payroll_engine.create_period(self.user, self.branch, 2026, 10)
        self.authenticate_client()
        self.assertEqual(self.client.get("/accounting/?start=2026-10-01&end=2026-10-31").status_code, 200)
        self.assertEqual(self.client.get("/reports/?family=expenses&category=fuel&start=2026-10-01&end=2026-10-31").status_code, 200)
        self.assertEqual(self.client.get("/payroll/").status_code, 200)
        exported = self.client.get("/accounting/export/pdf/?start=2026-10-01&end=2026-10-31")
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.content.startswith(b"%PDF"))


class CreditorsTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def credit_purchase(self, amount="60", due_date=None, external_reference="SUP-INV-001"):
        due_date = due_date or (timezone.localdate() + timedelta(days=14)).isoformat()
        return s.post_trade(
            self.user,
            self.branch,
            {
                "party": self.supplier.pk,
                "items": [{
                    "product": self.product.pk, "mode": "retail_unit",
                    "quantity": 1, "price": amount,
                }],
                "payments": [],
                "due_date": due_date,
                "external_reference": external_reference,
                "document_date": timezone.localdate().isoformat(),
                "note": "Supplier stock invoice",
            },
            uuid.uuid4(),
            "purchase",
        )

    def test_credit_purchase_automatically_becomes_supplier_payable(self):
        from . import creditors as creditor_service
        purchase = self.credit_purchase(amount="60")
        snapshot = creditor_service.supplier_account_snapshot(self.supplier)
        self.assertEqual(snapshot["outstanding"], Decimal("60"))
        self.assertEqual(snapshot["bill_count"], 1)
        self.assertEqual(snapshot["bill_rows"][0]["bill"], purchase)
        self.assertEqual(snapshot["bill_rows"][0]["source"], "Purchase")
        self.assertEqual(s.party_debt(self.supplier), Decimal("60"))

    def test_backdated_supplier_invoice_can_enter_as_already_overdue(self):
        from . import creditors as creditor_service
        invoice_date = timezone.localdate() - timedelta(days=45)
        due_date = timezone.localdate() - timedelta(days=15)
        purchase = s.post_trade(
            self.user, self.branch,
            {
                "party": self.supplier.pk,
                "items": [{"product": self.product.pk, "mode": "retail_unit", "quantity": 1, "price": "55"}],
                "payments": [],
                "document_date": invoice_date.isoformat(),
                "due_date": due_date.isoformat(),
                "external_reference": "HIST-55",
                "note": "Historical unpaid stock invoice",
            },
            uuid.uuid4(), "purchase",
        )
        snapshot = creditor_service.supplier_account_snapshot(self.supplier)
        self.assertEqual(purchase.document_date, invoice_date)
        self.assertEqual(purchase.due_date, due_date)
        self.assertEqual(snapshot["overdue"], Decimal("55"))
        self.assertEqual(snapshot["maximum_days_overdue"], 15)

    def test_direct_creditor_bill_creates_liability_without_stock_movement(self):
        from . import creditors as creditor_service
        before = Stock.objects.get(branch=self.branch, product=self.product).quantity
        bill = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "850.00",
                "document_date": timezone.localdate().isoformat(),
                "due_date": (timezone.localdate() + timedelta(days=7)).isoformat(),
                "external_reference": "RENT-OCT-2026", "category": "rent",
                "note": "October warehouse rent payable",
            },
            uuid.uuid4(),
        )
        self.assertEqual(bill.kind, "creditor_charge")
        self.assertEqual(s.balance(bill), Decimal("850"))
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, before)
        self.assertEqual(s.party_debt(self.supplier), Decimal("850"))

    def test_direct_bill_can_create_new_creditor_inline(self):
        from . import creditors as creditor_service
        bill = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "supplier_name": "New Service Vendor", "supplier_phone": "0249876543",
                "supplier_email": "vendor@example.test", "amount": "300",
                "document_date": timezone.localdate().isoformat(),
                "due_date": (timezone.localdate() + timedelta(days=10)).isoformat(),
                "external_reference": "NSV-001", "category": "professional",
                "note": "Professional service payable",
            },
            uuid.uuid4(),
        )
        self.assertEqual(bill.party.kind, "supplier")
        self.assertEqual(bill.party.name, "New Service Vendor")
        self.assertEqual(s.party_debt(bill.party), Decimal("300"))

    def test_duplicate_supplier_reference_is_blocked_across_purchase_and_direct_bill(self):
        from . import creditors as creditor_service
        self.credit_purchase(external_reference="VENDOR-77")
        with self.assertRaises(ValidationError):
            creditor_service.post_creditor_bill(
                self.user, self.branch,
                {
                    "party": str(self.supplier.pk), "amount": "120",
                    "document_date": timezone.localdate().isoformat(),
                    "due_date": (timezone.localdate() + timedelta(days=5)).isoformat(),
                    "external_reference": "vendor-77", "category": "professional",
                    "note": "Duplicate supplier bill reference",
                },
                uuid.uuid4(),
            )

    def test_account_payment_allocates_oldest_due_first_and_blocks_overpayment(self):
        from . import creditors as creditor_service
        old_bill = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "40",
                "document_date": (timezone.localdate() - timedelta(days=30)).isoformat(),
                "due_date": (timezone.localdate() - timedelta(days=10)).isoformat(),
                "external_reference": "OLD-40", "category": "utilities",
                "note": "Old utility creditor bill",
            },
            uuid.uuid4(),
        )
        purchase = self.credit_purchase(
            amount="60",
            due_date=(timezone.localdate() + timedelta(days=10)).isoformat(),
            external_reference="NEW-60",
        )
        payment = creditor_service.post_supplier_account_payment(
            self.user, self.branch,
            {"party": str(self.supplier.pk), "amount": "50", "method": "bank", "reference": "BANK-001"},
            uuid.uuid4(),
        )
        allocations = list(payment.allocations.order_by("pk"))
        self.assertEqual(allocations[0].invoice, old_bill)
        self.assertEqual(allocations[0].amount, Decimal("40"))
        self.assertEqual(allocations[1].invoice, purchase)
        self.assertEqual(allocations[1].amount, Decimal("10"))
        self.assertEqual(s.balance(old_bill), Decimal("0"))
        self.assertEqual(s.balance(purchase), Decimal("50"))
        self.assertEqual(s.channel_totals(self.branch, timezone.localdate())["bank"], Decimal("-50"))
        with self.assertRaises(ValidationError):
            creditor_service.post_supplier_account_payment(
                self.user, self.branch,
                {"party": str(self.supplier.pk), "amount": "51", "method": "bank"},
                uuid.uuid4(),
            )

    def test_selected_bill_payment_only_allocates_selected_bill(self):
        from . import creditors as creditor_service
        first = self.credit_purchase(amount="25", external_reference="ONE-25")
        second = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "35",
                "document_date": timezone.localdate().isoformat(),
                "due_date": (timezone.localdate() + timedelta(days=5)).isoformat(),
                "external_reference": "TWO-35", "category": "maintenance",
                "note": "Workshop repair payable",
            },
            uuid.uuid4(),
        )
        payment = creditor_service.post_supplier_account_payment(
            self.user, self.branch,
            {"party": str(self.supplier.pk), "invoice": str(second.pk), "amount": "20", "method": "cash"},
            uuid.uuid4(),
        )
        self.assertEqual(payment.allocations.count(), 1)
        self.assertEqual(payment.allocations.get().invoice, second)
        self.assertEqual(s.balance(first), Decimal("25"))
        self.assertEqual(s.balance(second), Decimal("15"))

    def test_direct_bill_correction_requires_payment_reversal_first(self):
        from . import creditors as creditor_service
        bill = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "90",
                "document_date": timezone.localdate().isoformat(),
                "due_date": (timezone.localdate() + timedelta(days=2)).isoformat(),
                "external_reference": "ERR-90", "category": "other",
                "note": "Creditor bill entered incorrectly",
            },
            uuid.uuid4(),
        )
        request = s.request_correction(self.user, self.branch, bill.pk, "Wrong creditor bill amount entered")
        s.review_correction(self.reviewer, self.branch, request.pk, True)
        bill.refresh_from_db()
        self.assertEqual(s.balance(bill), Decimal("0"))
        self.assertEqual(s.party_debt(self.supplier), Decimal("0"))

        paid_bill = creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "50",
                "document_date": timezone.localdate().isoformat(),
                "due_date": (timezone.localdate() + timedelta(days=2)).isoformat(),
                "external_reference": "PAID-50", "category": "other",
                "note": "Partially settled creditor bill",
            },
            uuid.uuid4(),
        )
        creditor_service.post_supplier_account_payment(
            self.user, self.branch,
            {"party": str(self.supplier.pk), "invoice": str(paid_bill.pk), "amount": "10", "method": "cash"},
            uuid.uuid4(),
        )
        with self.assertRaises(ValidationError):
            s.request_correction(self.user, self.branch, paid_bill.pk, "Need to reverse this paid creditor bill")

    def test_due_date_filter_scopes_creditor_totals_to_matching_bills(self):
        from . import creditors as creditor_service
        today = timezone.localdate()
        creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "40",
                "document_date": today.isoformat(),
                "due_date": (today + timedelta(days=3)).isoformat(),
                "external_reference": "DUE-3", "category": "utilities",
                "note": "Utility bill due soon",
            },
            uuid.uuid4(),
        )
        creditor_service.post_creditor_bill(
            self.user, self.branch,
            {
                "party": str(self.supplier.pk), "amount": "90",
                "document_date": today.isoformat(),
                "due_date": (today + timedelta(days=20)).isoformat(),
                "external_reference": "DUE-20", "category": "maintenance",
                "note": "Maintenance bill due later",
            },
            uuid.uuid4(),
        )
        self.authenticate_client()
        response = self.client.get("/creditors/", {
            "due_from": (today + timedelta(days=1)).isoformat(),
            "due_to": (today + timedelta(days=7)).isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["overview"]["total_payables"], Decimal("40"))
        self.assertEqual(response.context["rows"][0]["bill_count"], 1)
        self.assertEqual(response.context["rows"][0]["outstanding"], Decimal("40"))

    def test_creditor_search_finds_supplier_invoice_reference(self):
        from . import creditors as creditor_service
        self.credit_purchase(amount="75", external_reference="LOOKUP-AP-75")
        rows = creditor_service.creditor_accounts(self.branch, "lookup-ap-75")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["party"], self.supplier)

    def test_creditor_page_supplier_search_print_and_exports(self):
        self.credit_purchase(amount="75", external_reference="SEARCH-75")
        self.authenticate_client()
        page = self.client.get("/creditors/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, self.supplier.name)
        purchase_page = self.client.get("/purchasing/")
        self.assertContains(purchase_page, 'id="supplier-search"')
        self.assertContains(purchase_page, 'id="purchase-reference"')
        search = self.client.get("/api/suppliers/", {"q": "Supplier"})
        self.assertEqual(search.status_code, 200)
        self.assertEqual(search.json()["suppliers"][0]["outstanding"], "75.00")
        ref_search = self.client.get("/creditors/", {"q": "SEARCH-75"})
        self.assertContains(ref_search, self.supplier.name)
        for format in ("pdf", "xlsx", "docx", "csv"):
            response = self.client.get(f"/creditors/export/{format}/")
            self.assertEqual(response.status_code, 200)
        printable = self.client.get(
            f"/documents/{Document.objects.get(external_reference='SEARCH-75').pk}/pdf/a4/"
        )
        self.assertEqual(printable.status_code, 200)
        self.assertTrue(printable.content.startswith(b"%PDF"))




class ControlCentreIntelligenceTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def staff_with(self, username, codename):
        user = User.objects.create_user(username, password="control-test-password-long-enough")
        user.user_permissions.add(Permission.objects.get(codename=codename))
        user.access.branches.add(self.branch)
        return user

    def test_customer_return_waits_for_approval_without_direct_privilege(self):
        from . import returns as return_controls
        from .models import CustomerReturnRequest
        sale = self.sale(2)
        source = sale.lines.get()
        cashier = self.staff_with("return-cashier", "operate_sales")
        before = Stock.objects.get(branch=self.branch, product=self.product).quantity
        item, direct = return_controls.create_customer_return(
            cashier, self.branch, sale,
            [{"line": str(source.pk), "quantity": "1", "disposition": "sellable"}],
            "Customer brought back one item", "cash",
        )
        self.assertFalse(direct)
        self.assertEqual(item.status, "requested")
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, before)
        self.assertFalse(Document.objects.filter(kind="return", original=sale).exists())
        return_controls.execute_customer_return(self.reviewer, self.branch, item.pk)
        item = CustomerReturnRequest.objects.get(pk=item.pk)
        self.assertEqual(item.status, "approved")
        self.assertIsNotNone(item.posted_id)
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, before + 1)

    def test_return_privilege_posts_directly_and_quarantine_never_becomes_sellable(self):
        from . import returns as return_controls
        from .models import QuarantineItem, ReturnPrivilege
        sale = self.sale()
        source = sale.lines.get()
        cashier = self.staff_with("direct-return-user", "operate_sales")
        ReturnPrivilege.objects.create(
            branch=self.branch, user=cashier, customer_returns=True, granted_by=self.user
        )
        before = Stock.objects.get(branch=self.branch, product=self.product).quantity
        item, direct = return_controls.create_customer_return(
            cashier, self.branch, sale,
            [{"line": str(source.pk), "quantity": "1", "disposition": "quarantine"}],
            "Item returned damaged and unsafe to resell", "cash",
        )
        self.assertTrue(direct)
        item.refresh_from_db()
        self.assertEqual(item.status, "approved")
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, before)
        held = QuarantineItem.objects.get(reason__contains=item.posted.reference)
        self.assertEqual(held.status, "held")
        self.assertEqual(held.quantity, 1)

    def test_supplier_return_privilege_posts_without_separate_approval(self):
        from . import inventory_exceptions
        from .models import ReturnPrivilege, SupplierReturn
        purchase = s.post_trade(
            self.user, self.branch,
            {
                "party": self.supplier.pk,
                "items": [{"product": self.product.pk, "mode": "retail_unit", "quantity": 2, "price": "20"}],
                "payments": [{"method": "cash", "amount": "40"}],
                "document_date": timezone.localdate().isoformat(),
                "external_reference": "DIRECT-SR-1",
            },
            uuid.uuid4(), "purchase",
        )
        keeper = self.staff_with("direct-supplier-return", "operate_inventory")
        ReturnPrivilege.objects.create(
            branch=self.branch, user=keeper, supplier_returns=True, granted_by=self.user
        )
        before = Stock.objects.get(branch=self.branch, product=self.product).quantity
        item = inventory_exceptions.request_supplier_return(
            keeper, self.branch, purchase.lines.get().pk, 1,
            "Supplier accepted one incorrect item back", "cash", direct=True,
        )
        item = SupplierReturn.objects.get(pk=item.pk)
        self.assertEqual(item.status, "approved")
        self.assertIsNotNone(item.posted_id)
        self.assertEqual(Stock.objects.get(branch=self.branch, product=self.product).quantity, before - 1)

    def test_double_entry_trial_balance_and_manual_journal_control(self):
        from . import accounting_engine
        self.sale()
        today = timezone.localdate()
        tb = accounting_engine.trial_balance(self.branch, today, today)
        self.assertEqual(
            sum((row["debit"] for row in tb), Decimal("0")),
            sum((row["credit"] for row in tb), Decimal("0")),
        )
        statements = accounting_engine.statements(self.branch, today, today)
        self.assertEqual(statements["balance_check"], Decimal("0.00"))
        with self.assertRaises(ValidationError):
            accounting_engine.create_manual_journal(
                self.user, self.branch, today.isoformat(), "Unbalanced test journal",
                [{"account_code": "1000", "debit": "10", "credit": "0"},
                 {"account_code": "3000", "debit": "0", "credit": "9"}],
            )
        journal = accounting_engine.create_manual_journal(
            self.user, self.branch, today.isoformat(), "Owner capital introduced",
            [{"account_code": "1000", "debit": "100", "credit": "0", "description": "Cash introduced"},
             {"account_code": "3000", "debit": "0", "credit": "100", "description": "Owner capital"}],
        )
        self.assertEqual(journal.status, "posted")

    def test_audit_events_are_hash_linked(self):
        from .models import Audit
        first = s.audit(self.user, self.branch, "test.first", "A", {"value": 1})
        second = s.audit(self.user, self.branch, "test.second", "B", {"value": 2})
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(len(first.event_hash), 64)
        self.assertEqual(second.previous_hash, first.event_hash)
        self.assertNotEqual(first.event_hash, second.event_hash)
        self.assertEqual(Audit.objects.filter(branch=self.branch, event_hash="").count(), 0)

    def test_approval_center_return_privileges_intelligence_and_accounting_pages(self):
        self.sale()
        self.authenticate_client()
        self.assertEqual(self.client.get("/approvals/").status_code, 200)
        self.assertEqual(self.client.get("/api/approvals/summary/").status_code, 200)
        self.assertEqual(self.client.get("/settings/return-privileges/").status_code, 200)
        intelligence = self.client.get("/reports/?family=sales")
        self.assertEqual(intelligence.status_code, 200)
        self.assertContains(intelligence, "BUSINESS INTELLIGENCE COMMAND CENTRE")
        accounting = self.client.get("/accounting/?view=trial")
        self.assertEqual(accounting.status_code, 200)
        self.assertContains(accounting, "Trial balance")
        audit = self.client.get("/audit/")
        self.assertEqual(audit.status_code, 200)
        self.assertContains(audit, "AUDIT INTELLIGENCE")

    def test_stock_operations_are_retired_in_production(self):
        from django.test import override_settings
        self.authenticate_client()
        with override_settings(DEBUG=False):
            response = self.client.get("/operations/")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/inventory/")



class CorrectionTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()

    def test_payment_reversal_restores_debt_and_channel(self):
        sale = self.sale(payments=[],party=self.customer.pk,due_date=timezone.localdate().isoformat())
        payment = s.post_payment(self.user,self.branch,{"invoice":str(sale.pk),"amount":"50","method":"momo"},uuid.uuid4())
        request = s.request_correction(self.user,self.branch,payment.pk,"Wrong payment reference recorded")
        with self.assertRaises(ValidationError):
            s.review_correction(self.user,self.branch,request.pk,True)
        s.review_correction(self.reviewer,self.branch,request.pk,True)
        self.assertEqual(s.balance(sale),Decimal("50"))
        self.assertEqual(s.channel_totals(self.branch,timezone.localdate())["momo"],Decimal("0"))
        with self.assertRaises(ValidationError):
            s.review_correction(self.reviewer,self.branch,request.pk,True)

    def test_void_sale_preserves_original_and_restores_stock(self):
        sale = self.sale(2)
        request = s.request_correction(self.user,self.branch,sale.pk,"Duplicate counter sale entered")
        s.review_correction(self.reviewer,self.branch,request.pk,True)
        sale.refresh_from_db()
        self.assertEqual(sale.total,Decimal("100"))
        self.assertEqual(Stock.objects.get(branch=self.branch,product=self.product).quantity,240)
        self.assertEqual(s.channel_totals(self.branch,timezone.localdate())["cash"],Decimal("0"))

    def test_expense_reversal_preserves_evidence(self):
        expense = s.post_expense(self.user,self.branch,{"amount":"20","method":"bank","note":"Office supply expense"},uuid.uuid4())
        request = s.request_correction(self.user,self.branch,expense.pk,"Incorrect supplier expense posted")
        result = s.review_correction(self.reviewer,self.branch,request.pk,True)
        self.assertEqual(result.posted.original,expense)
        self.assertEqual(s.channel_totals(self.branch,timezone.localdate())["bank"],Decimal("0"))


class ReportingTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()
    def test_report_families_and_exports(self):
        self.sale(payments=[],party=self.customer.pk,due_date=timezone.localdate().isoformat())
        self.authenticate_client()
        for family in ("register","sales","inventory","aging","creditors"):
            self.assertEqual(self.client.get("/reports/?family="+family).status_code,200)
            self.assertEqual(self.client.get("/reports/export/pdf/?family="+family).status_code,200)
    def test_profit_uses_original_standard_cost(self):
        from .reporting import build_report
        self.sale(2)
        self.product.cost = Decimal("30")
        self.product.save()
        rows,_ = build_report(self.branch,timezone.localdate(),timezone.localdate(),"sales")
        self.assertEqual(rows[0]["profit"],Decimal("60"))
    def test_global_search_scopes_transactions(self):
        doc = self.sale()
        self.authenticate_client()
        response = self.client.get("/search/",{"q":doc.reference})
        self.assertContains(response,doc.reference)


class OwnerSecurityTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()
    def test_company_manager_has_no_authenticator_gate(self):
        owner = User.objects.create_user("business-owner",password="business-owner-password")
        owner.user_permissions.add(Permission.objects.get(codename="manage_company"))
        owner.access.branches.add(self.branch)
        self.client.post("/login/",{"username":"business-owner","password":"business-owner-password"})
        self.assertEqual(self.client.get("/").url,"/inventory/")
    def test_cashier_cannot_search_supplier_contacts(self):
        cashier = User.objects.create_user("limited-cashier",password="limited-cashier-password")
        cashier.user_permissions.add(Permission.objects.get(codename="operate_sales"))
        cashier.access.branches.add(self.branch)
        self.authenticate_client(cashier)
        self.assertNotContains(self.client.get("/search/",{"q":"Supplier"}),self.supplier.phone)



class AdministrationAndExportTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()
        self.authenticate_client()

    def test_admin_route_is_business_centre_not_django_index(self):
        response = self.client.get("/admin/")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/administration/")
        page = self.client.get("/administration/")
        self.assertContains(page, "Run KOFAD without entering the technical backend")
        self.assertNotContains(page, "Authentication and Authorization")

    def test_owner_can_create_cashier_in_branded_admin(self):
        role = Group.objects.create(name="Cashier test")
        role.permissions.add(Permission.objects.get(codename="operate_sales"))
        response = self.client.post("/administration/users/new/", {
            "username": "counter-one",
            "first_name": "Counter",
            "last_name": "One",
            "role": str(role.pk),
            "password": "Strong-counter-password-2026!",
            "active": "on",
            "branches": [str(self.branch.pk)],
            "recovery_phone": "0241234567",
        })
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="counter-one")
        self.assertTrue(user.groups.filter(pk=role.pk).exists())
        self.assertTrue(user.access.branches.filter(pk=self.branch.pk).exists())
        self.assertEqual(user.access.recovery_phone, "+233241234567")
        self.assertFalse(user.is_staff)

    def test_role_editor_changes_only_selected_role_members_sessions(self):
        role = Group.objects.create(name="Counter control")
        role.permissions.add(Permission.objects.get(codename="operate_sales"))
        member = User.objects.create_user("role-member", password="role-member-strong-password")
        member.groups.add(role)
        member.access.branches.add(self.branch)
        outsider = User.objects.create_user("role-outsider", password="role-outsider-strong-password")
        outsider.access.branches.add(self.branch)
        member.access.refresh_from_db()
        outsider.access.refresh_from_db()
        before_member = member.access.session_version
        before_outsider = outsider.access.session_version
        self.client.post(f"/administration/roles/{role.pk}/", {
            "name": "Counter control",
            "permissions": [
                str(Permission.objects.get(codename="operate_sales").pk),
                str(Permission.objects.get(codename="add_party").pk),
            ],
        })
        member.access.refresh_from_db()
        outsider.access.refresh_from_db()
        self.assertGreater(member.access.session_version, before_member)
        self.assertEqual(outsider.access.session_version, before_outsider)

    def test_export_centre_produces_business_datasets_in_all_formats(self):
        self.sale()
        from openpyxl import load_workbook
        from docx import Document as WordDocument
        for dataset in ("customers", "creditors", "inventory", "sales", "transactions", "movements", "audit", "staff"):
            with self.subTest(dataset=dataset):
                response = self.client.get(f"/exports/download/xlsx/?dataset={dataset}")
                self.assertEqual(response.status_code, 200)
                self.assertTrue(load_workbook(BytesIO(response.content)).active.max_row >= 2)
        for format in ("pdf", "docx", "csv"):
            response = self.client.get(f"/exports/download/{format}/?dataset=customers")
            self.assertEqual(response.status_code, 200)
            if format == "pdf":
                self.assertTrue(response.content.startswith(b"%PDF"))
            if format == "docx":
                self.assertGreaterEqual(len(WordDocument(BytesIO(response.content)).tables), 1)

    def test_single_store_hides_location_switcher(self):
        self.other.delete()
        response = self.client.get("/")
        self.assertNotContains(response, 'id="branch-select"')
        self.assertNotContains(response, ">Switch<")
        self.assertContains(response, self.branch.name)
