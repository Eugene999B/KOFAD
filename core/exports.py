import csv
import io
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.utils import timezone

COLUMNS = ["reference", "date", "kind", "party", "total", "paid", "balance"]
CHARCOAL, COPPER, CREAM, LIGHT, WHITE = "142B3B", "96603B", "F8F5EE", "E8EEF0", "FFFFFF"


def safe(value):
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def display(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    return str(value)


def excel_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if isinstance(value, (int, float, date, datetime)):
        return value
    return safe(value)


def shade_word(cell, colour):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), colour)


def export(
    rows, format, title, company, columns=None, filename="kofad-transactions",
    sheet_name="Transactions", metadata=None, summary=None, notes=None,
):
    columns = columns or [(key, key.title()) for key in COLUMNS]
    headers = [label for _, label in columns]
    raw = [[row.get(key, "") for key, _ in columns] for row in rows]
    shown = [[display(value) for value in row] for row in raw]
    metadata = [(str(k), display(v)) for k, v in (metadata or {}).items() if v not in (None, "")]
    summary = [(str(k), v) for k, v in (summary or {}).items()]
    notes = [str(item) for item in (notes or []) if item]
    output = io.BytesIO()

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = (sheet_name or "Export")[:31]
        book.properties.title = title
        book.properties.creator = company.name
        width = max(1, len(headers))
        row_no = 1

        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=width)
        cell = sheet.cell(row_no, 1, company.name)
        cell.font = Font(size=18, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=CHARCOAL)
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row_no].height = 32
        row_no += 1

        sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=width)
        cell = sheet.cell(row_no, 1, title)
        cell.font = Font(size=12, bold=True, color=CHARCOAL)
        cell.fill = PatternFill("solid", fgColor=CREAM)
        sheet.row_dimensions[row_no].height = 24
        row_no += 1

        for key, value in metadata:
            sheet.cell(row_no, 1, key).font = Font(bold=True, color="6B6B6B")
            if width > 1:
                sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=width)
                sheet.cell(row_no, 2, excel_value(value))
            row_no += 1

        if summary:
            row_no += 1
            card_cols = min(4, width)
            start = row_no
            for index, (key, value) in enumerate(summary):
                col = index % card_cols + 1
                card_row = start + (index // card_cols) * 2
                sheet.cell(card_row, col, key)
                sheet.cell(card_row + 1, col, excel_value(value))
                sheet.cell(card_row, col).fill = PatternFill("solid", fgColor=LIGHT)
                sheet.cell(card_row, col).font = Font(size=9, bold=True, color="6B6B6B")
                sheet.cell(card_row + 1, col).font = Font(size=12, bold=True, color=CHARCOAL)
                if isinstance(value, (Decimal, int, float)) and not isinstance(value, bool):
                    sheet.cell(card_row + 1, col).number_format = '#,##0' if isinstance(value, int) else '#,##0.00'
            row_no = start + ((len(summary) - 1) // card_cols + 1) * 2 + 1

        header_row = row_no
        for index, header in enumerate(headers, 1):
            cell = sheet.cell(header_row, index, header)
            cell.fill = PatternFill("solid", fgColor=COPPER)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[header_row].height = 28

        border = Border(bottom=Side(style="thin", color="DDD8D0"))
        for row_index, record in enumerate(raw, header_row + 1):
            for col_index, value in enumerate(record, 1):
                cell = sheet.cell(row_index, col_index, excel_value(value))
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="FAF8F4")
                if isinstance(value, (Decimal, int, float)) and not isinstance(value, bool):
                    cell.number_format = '#,##0' if isinstance(value, int) else '#,##0.00'

        if headers:
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(width)}{max(header_row, sheet.max_row)}"
        for index, header in enumerate(headers, 1):
            lengths = [len(str(header))]
            lengths.extend(len(row[index - 1]) for row in shown[:200] if index - 1 < len(row))
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max(lengths) + 3, 12), 42)
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = "$%s:$%s" % (header_row, header_row)
        sheet.page_setup.orientation = "landscape" if width > 6 else "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddFooter.center.text = f"{company.name} · KOFAD generated export"
        sheet.oddFooter.right.text = "Page &P of &N"
        if notes:
            notes_sheet = book.create_sheet("Report notes")
            notes_sheet.column_dimensions["A"].width = 100
            for note in notes:
                notes_sheet.append([safe(note)])
                notes_sheet.cell(notes_sheet.max_row, 1).alignment = Alignment(wrap_text=True)
        book.save(output)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    elif format == "docx":
        from docx import Document as WordDocument
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor

        doc = WordDocument()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE if len(headers) > 6 else WD_ORIENT.PORTRAIT
        if section.orientation == WD_ORIENT.LANDSCAPE:
            section.page_width, section.page_height = Inches(11.7), Inches(8.3)
        else:
            section.page_width, section.page_height = Inches(8.3), Inches(11.7)
        section.left_margin = section.right_margin = Inches(.55)
        section.top_margin = section.bottom_margin = Inches(.55)
        doc.styles["Normal"].font.name = "Aptos"
        doc.styles["Normal"].font.size = Pt(9)

        p = doc.add_paragraph()
        run = p.add_run(company.name)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(23, 23, 23)
        p.paragraph_format.space_after = Pt(1)
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(184, 115, 51)
        p.paragraph_format.space_after = Pt(8)

        if metadata:
            meta_table = doc.add_table(rows=0, cols=2)
            for key, value in metadata:
                cells = meta_table.add_row().cells
                cells[0].text, cells[1].text = key, value
                shade_word(cells[0], CREAM)
                cells[0].paragraphs[0].runs[0].bold = True
            doc.add_paragraph()

        if summary:
            summary_table = doc.add_table(rows=2, cols=min(4, len(summary)))
            for index, (key, value) in enumerate(summary[:4]):
                summary_table.cell(0, index).text = key
                summary_table.cell(1, index).text = display(value)
                shade_word(summary_table.cell(0, index), CHARCOAL)
                shade_word(summary_table.cell(1, index), CREAM)
                for run in summary_table.cell(0, index).paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                for run in summary_table.cell(1, index).paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(11)
            doc.add_paragraph()

        from docx.oxml import OxmlElement
        groups = [list(range(len(headers)))] if len(headers) <= 9 else [
            [0] + list(range(index, min(index + 7, len(headers)))) for index in range(1, len(headers), 7)
        ]
        for group_index, group in enumerate(groups):
            if group_index:
                doc.add_page_break()
                doc.add_heading(title + " - continued fields", level=2)
            group_headers = [headers[index] for index in group]
            group_records = [[record[index] for index in group] for record in shown]
            table = doc.add_table(rows=1, cols=len(group_headers))
            table.style = "Table Grid"
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell, header in zip(table.rows[0].cells, group_headers):
                cell.text = header
                shade_word(cell, COPPER)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(8)
            for index, record in enumerate(group_records):
                cells = table.add_row().cells
                for column_index, (cell, value) in enumerate(zip(cells, record)):
                    cell.text = value
                    if isinstance(raw[index][group[column_index]], (int, float, Decimal)):
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if index % 2:
                        shade_word(cell, "FAF8F4")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.size = Pt(9 if len(headers) <= 7 else 8)

        for note in notes:
            p = doc.add_paragraph()
            run = p.add_run(note)
            run.italic = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(107, 107, 107)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(f"{company.name} · Generated by KOFAD").font.size = Pt(7)
        doc.save(output)
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    elif format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape, portrait
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        from .brand_art import draw_mark, print_fonts
        regular, bold = print_fonts()
        page_size = landscape(A4) if len(headers) > 6 else portrait(A4)
        styles = getSampleStyleSheet()
        body_style = ParagraphStyle("KofadBody", parent=styles["BodyText"], fontName=regular, fontSize=8.3, leading=11.5)
        number_style = ParagraphStyle("KofadNumber", parent=body_style, alignment=2)
        head_style = ParagraphStyle("KofadHead", parent=body_style, fontName=bold, textColor=colors.white, leading=11)
        small_style = ParagraphStyle("KofadSmall", parent=body_style, fontSize=8, leading=11, textColor=colors.HexColor("#526271"))
        title_style = ParagraphStyle("KofadTitle", parent=styles["Heading1"], fontName=bold, fontSize=20, leading=25, textColor=colors.HexColor("#171717"))

        def chrome(pdf, document):
            page_w, page_h = page_size
            pdf.saveState()
            pdf.setFillColor(colors.HexColor("#171717"))
            pdf.rect(0, page_h - 14 * mm, page_w, 14 * mm, fill=1, stroke=0)
            pdf.setFillColor(colors.white)
            pdf.roundRect(12 * mm, page_h - 12 * mm, 10 * mm, 10 * mm, 1.5 * mm, fill=1, stroke=0)
            draw_mark(pdf, 13 * mm, page_h - 11 * mm, 8 * mm)
            pdf.setFillColor(colors.HexColor("#B87333"))
            pdf.rect(0, page_h - 10 * mm, 3 * mm, 10 * mm, fill=1, stroke=0)
            pdf.setFillColor(colors.white)
            pdf.setFont("Helvetica-Bold", 9)
            pdf.drawString(26 * mm, page_h - 8.4 * mm, "KOFAD IMPEX ENTERPRISE")
            pdf.setFillColor(colors.HexColor("#777777"))
            pdf.setFont("Helvetica", 6.5)
            pdf.drawString(12 * mm, 7 * mm, "Generated by KOFAD · Controlled business record")
            pdf.drawRightString(page_w - 12 * mm, 7 * mm, f"Page {document.page}")
            pdf.restoreState()

        story = [Paragraph(escape(title), title_style), Spacer(1, 4)]
        if metadata:
            meta_rows = [[Paragraph(f"<b>{escape(key)}</b>", small_style), Paragraph(escape(value), small_style)] for key, value in metadata]
            meta_table = Table(meta_rows, colWidths=[36 * mm, None])
            meta_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F6F2EA")),
                ("LINEBELOW", (0, 0), (-1, -1), .25, colors.HexColor("#DDD8D0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([meta_table, Spacer(1, 7)])
        if summary:
            summary_rows = [[Paragraph(escape(key), small_style), Paragraph(f"<b>{escape(display(value))}</b>", body_style)] for key, value in summary]
            summary_table = Table(summary_rows, colWidths=[45 * mm, 45 * mm])
            summary_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F6F2EA")),
                ("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#DDD8D0")),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([summary_table, Spacer(1, 7)])

        # Wide registers continue in numbered field groups rather than shrinking
        # every column to unreadable type. Row numbers preserve cross-page identity.
        groups = [list(range(len(headers)))] if len(headers) <= 9 else [
            [0] + list(range(index, min(index + 7, len(headers)))) for index in range(1, len(headers), 7)
        ]
        available_width = page_size[0] - 24 * mm
        for group_index, group in enumerate(groups):
            if group_index:
                story.extend([PageBreak(), Paragraph(escape(title) + " - continued fields", title_style), Spacer(1, 10)])
            group_headers = [headers[index] for index in group]
            if len(groups) > 1:
                group_headers = ["Row"] + group_headers
            cells = [[Paragraph(escape(str(value)), head_style) for value in group_headers]]
            for row_number, record in enumerate(shown, 1):
                values = [record[index] for index in group]
                if len(groups) > 1:
                    values = [str(row_number)] + values
                raw_values = [raw[row_number - 1][index] for index in group]
                if len(groups) > 1:
                    raw_values = [row_number] + raw_values
                cells.append([Paragraph(escape(value).replace("\n", "<br/>"),
                    number_style if isinstance(raw_value, (int, float, Decimal)) else body_style)
                    for value, raw_value in zip(values, raw_values)])
            weights = []
            for index in group:
                key = columns[index][0]
                weights.append(2.1 if any(token in key for token in ("name", "customer", "product", "contact", "note", "reason", "evidence")) else 1.25)
            if len(groups) > 1:
                weights.insert(0, .5)
            widths = [available_width * weight / sum(weights) for weight in weights]
            table = Table(cells, repeatRows=1, colWidths=widths, hAlign="LEFT")
            commands = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#142B3B")),
                ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.HexColor("#B87333")),
                ("LINEBELOW", (0, 1), (-1, -1), .25, colors.HexColor("#DCE4E8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
            for row_index in range(2, len(cells), 2):
                commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F3F6F7")))
            table.setStyle(TableStyle(commands))
            story.append(table)
            if not shown:
                story.extend([Spacer(1, 10), Paragraph("No records match this report.", body_style)])
        if notes:
            story.append(Spacer(1, 8))
            story.extend(Paragraph(escape(note), small_style) for note in notes)

        SimpleDocTemplate(
            output, pagesize=page_size, leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=24 * mm, bottomMargin=14 * mm, title=title, author=company.name,
        ).build(story, onFirstPage=chrome, onLaterPages=chrome)
        content_type = "application/pdf"

    elif format == "csv":
        text = io.StringIO()
        writer = csv.writer(text)
        writer.writerow(headers)
        # Preserve genuine signed numeric values. Escape only untrusted text,
        # otherwise negative amounts become apostrophe-prefixed strings in CSV.
        writer.writerows([
            [display(value) if isinstance(value, (Decimal, int, float)) and not isinstance(value, bool)
             else safe(display(value)) for value in record]
            for record in raw
        ])
        output.write(text.getvalue().encode("utf-8-sig"))
        content_type = "text/csv"
    else:
        raise ValidationError("Unsupported export format.")

    response = HttpResponse(output.getvalue(), content_type=content_type)
    safe_filename = "".join(ch for ch in filename.lower() if ch.isalnum() or ch in "-_") or "kofad-export"
    response["Content-Disposition"] = f'attachment; filename="{safe_filename}.{format}"'
    return response
