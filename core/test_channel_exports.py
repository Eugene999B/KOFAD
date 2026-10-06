import io
import json
import uuid
from datetime import timedelta
from decimal import Decimal
from unittest.mock import Mock, patch

import requests
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from .export_views import select_columns
from .exports import export
from .models import CommunicationSettings, Company, Message, WhatsAppWebhookEvent
from .sms.service import create_draft
from .tests import Fixtures
from .whatsapp import _reconcile_status
from .whatsapp_delivery import process_whatsapp_queue, queue_whatsapp, send_whatsapp


class ChannelAndExportTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()
        self.customer.phone = "+233241234567"
        self.customer.consent = True
        self.customer.save()
        self.authenticate_client()

    @override_settings(SMS_ENABLED=True)
    @patch("core.sms.providers.Arkesel.submit_many")
    def test_explicit_off_prevents_automatic_receipt_and_repeat(self, submit):
        CommunicationSettings.objects.update_or_create(pk=1, defaults={"sale_receipt_mode": "send", "whatsapp_sale_receipt_mode": "send"})
        payload = {"party": self.customer.pk, "customer_consent": False,
                   "send_sms": False, "send_whatsapp": False,
                   "items": [{"product": self.product.pk, "mode": "retail_unit", "quantity": 1}],
                   "payments": [{"method": "cash", "amount": "50"}]}
        key = str(uuid.uuid4())
        with self.captureOnCommitCallbacks(execute=True):
            first = self.client.post("/api/trades/", json.dumps(payload), content_type="application/json", HTTP_IDEMPOTENCY_KEY=key)
            repeat = self.client.post("/api/trades/", json.dumps(payload), content_type="application/json", HTTP_IDEMPOTENCY_KEY=key)
        self.assertEqual(first.status_code, 200)
        self.assertEqual(repeat.json()["document_id"], first.json()["document_id"])
        self.assertFalse(first.json()["sms_requested"])
        self.assertFalse(first.json()["whatsapp_requested"])
        self.assertFalse(Message.objects.exists())
        submit.assert_not_called()

    def test_customer_register_has_only_operational_columns(self):
        columns = [(key, key.title()) for key in ("name", "email", "phone", "address", "messages", "outstanding")]
        self.assertEqual([key for key, _ in select_columns("customers", columns)], ["name", "phone", "outstanding"])

    def test_excel_handles_aware_dates_numeric_values_and_formula_text(self):
        now = timezone.now()
        response = export([{"name": "=1+1", "created": now, "balance": Decimal("21.50")}],
                          "xlsx", "Customer register", Company.objects.get(),
                          [("name", "Customer"), ("created", "Joined"), ("balance", "Balance")])
        book = load_workbook(io.BytesIO(response.content))
        rows = list(book.active.values)
        row = next(row for row in rows if row[0] == "'=1+1")
        self.assertIsNone(row[1].tzinfo)
        self.assertEqual(row[2], 21.5)
        self.assertIsNotNone(book.active.auto_filter.ref)

    def test_wide_pdf_and_word_preserve_all_columns(self):
        columns = [(f"field{i}", f"Field {i}") for i in range(17)]
        row = {key: "Long operational detail with wrapped text" for key, _ in columns}
        for kind in ("pdf", "docx"):
            result = export([row], kind, "Wide register", Company.objects.get(), columns,
                            metadata={"Location": "Kumasi"}, summary={"Rows": 1})
            self.assertGreater(len(result.content), 1000)
            if kind == "docx":
                from docx import Document
                document = Document(io.BytesIO(result.content))
                text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                self.assertIn("Field 16", text)

    def test_receipt_api_rejects_non_object_payload(self):
        sale = self.sale(send_sms=False, send_whatsapp=False)
        response = self.client.post(f"/api/documents/{sale.pk}/send-sms/", "[]", content_type="application/json")
        self.assertEqual(response.status_code, 400)


@override_settings(WHATSAPP_ENABLED=True, WHATSAPP_ACCESS_TOKEN="isolated-test-token",
                   WHATSAPP_PHONE_NUMBER_ID="12345", WHATSAPP_APP_SECRET="isolated-test-secret",
                   WHATSAPP_GRAPH_VERSION="v23.0", WHATSAPP_TIMEOUT_SECONDS=10)
class WhatsAppDeliveryTests(Fixtures, TestCase):
    def setUp(self):
        self.setup_data()
        self.customer.phone = "+233241234567"
        self.customer.consent = True
        self.customer.save()
        CommunicationSettings.objects.update_or_create(pk=1, defaults={"whatsapp_template_name": "kofad_notification", "whatsapp_template_language": "en"})
        self.message = create_draft(self.user, self.branch, self.customer, "Your receipt is ready.", channel="whatsapp")

    @patch("core.whatsapp_delivery.requests.post")
    def test_template_accepted_then_delivery_webhook_and_no_duplicate(self, post):
        post.return_value = Mock(status_code=200, ok=True, json=lambda: {"messages": [{"id": "wamid.test"}]})
        result = send_whatsapp(self.user, self.branch, self.message.pk, automatic=True)
        self.assertEqual(result.status, "accepted")
        payload = post.call_args.kwargs["json"]
        self.assertEqual(payload["type"], "template")
        self.assertEqual(payload["template"]["name"], "kofad_notification")
        _reconcile_status("wamid.test", "delivered", {})
        result = send_whatsapp(self.user, self.branch, self.message.pk, automatic=True)
        self.assertEqual(result.status, "delivered")
        post.assert_called_once()

    @patch("core.whatsapp_delivery.requests.post", side_effect=requests.Timeout)
    def test_uncertain_delivery_is_not_automatically_retried(self, post):
        for _ in range(2):
            result = send_whatsapp(self.user, self.branch, self.message.pk, automatic=True, retry=True)
            self.assertEqual(result.status, "unknown")
        post.assert_called_once()
        self.assertEqual(self.message.whatsapp_attempts.count(), 1)

    @patch("core.whatsapp_delivery.requests.post")
    def test_recent_arrival_of_old_inbound_does_not_open_freeform_window(self, post):
        post.return_value = Mock(status_code=200, ok=True, json=lambda: {"messages": [{"id": "wamid.old"}]})
        WhatsAppWebhookEvent.objects.create(fingerprint="old", event_type="message", wa_id="233241234567",
            phone_number_id="12345", payload={"timestamp": str(int((timezone.now() - timedelta(days=2)).timestamp()))})
        send_whatsapp(self.user, self.branch, self.message.pk)
        self.assertEqual(post.call_args.kwargs["json"]["type"], "template")

    @patch("core.whatsapp_delivery.requests.post")
    def test_actual_recent_inbound_allows_manual_text(self, post):
        post.return_value = Mock(status_code=200, ok=True, json=lambda: {"messages": [{"id": "wamid.recent"}]})
        WhatsAppWebhookEvent.objects.create(fingerprint="recent", event_type="message", wa_id="233241234567",
            phone_number_id="12345", payload={"timestamp": str(int(timezone.now().timestamp()))})
        send_whatsapp(self.user, self.branch, self.message.pk)
        self.assertEqual(post.call_args.kwargs["json"]["type"], "text")

    @patch("core.whatsapp_delivery.requests.post")
    def test_queue_is_durable_and_worker_submits_once(self, post):
        post.return_value = Mock(status_code=200, ok=True, json=lambda: {"messages": [{"id": "wamid.queued"}]})
        for _ in range(2):
            result = queue_whatsapp(self.user, self.branch, self.message.pk, automatic=True)
            self.assertEqual(result.status, "queued")
        post.assert_not_called()
        self.assertEqual(process_whatsapp_queue(), 1)
        self.assertEqual(process_whatsapp_queue(), 0)
        self.message.refresh_from_db()
        self.assertEqual(self.message.status, "accepted")
        post.assert_called_once()
